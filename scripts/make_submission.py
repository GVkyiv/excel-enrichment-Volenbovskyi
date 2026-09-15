"""Готує файли на здачу: копії результатів, підписані прізвищем.

Умова курсу вимагає підписувати файли власним прізвищем. Робочі імена
файлів при цьому лишаються технічними, щоб їх не доводилось міняти в коді
і в ноутбуці.

Запуск: py scripts\\make_submission.py
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from scripts.validate_against_reference import (  # noqa: E402
    TOLERANCE,
    detect_alt_column,
    read_column,
)

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "Data" / "Output"
REFERENCE = ROOT / "Data" / "Reference"
SUBMISSION = ROOT / "Data" / "Submission"
SURNAME = "ВоленбовськийГВ"

FILES = ["capitals_enriched.xlsx", "mountains_enriched.xlsx"]

COMPARISON_SHEET = "звірка_з_еталоном"

# Колонка, за якою звіряємось, для кожного набору.
COLUMNS = {"capitals_enriched.xlsx": "distance", "mountains_enriched.xlsx": "height"}


def add_comparison_sheet(path: Path, reference_path: Path, column: str) -> None:
    """Додає в книгу аркуш зі звіркою з еталоном викладача.

    Потрібен на випадок, коли перевіряють самі файли, не заглядаючи в
    README: розбіжність на трьох рядках має бути пояснена там само, де її
    видно.
    """
    if not reference_path.exists():
        return
    workbook = load_workbook(path)
    if COMPARISON_SHEET in workbook.sheetnames:
        del workbook[COMPARISON_SHEET]
    sheet = workbook.create_sheet(COMPARISON_SHEET)

    alt_column = detect_alt_column(path, reference_path)
    headers = ["рядок", "еталон", "наше значення", "відхилення, %"]
    if alt_column:
        headers.append(alt_column)
    headers.append("вердикт")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    result = read_column(path, column)
    expected = read_column(path=reference_path, column=column)
    alternative = read_column(path, alt_column) if alt_column else {}

    matched = matched_any = checked = 0
    for row, reference_value in expected.items():
        if reference_value is None:
            continue
        checked += 1
        actual = result.get(row)
        try:
            deviation = (float(actual) - float(reference_value)) / float(reference_value)
        except (TypeError, ValueError):
            sheet.append([row, reference_value, actual, None, "не порівнюється"])
            continue
        in_range = abs(deviation) <= TOLERANCE
        matched += in_range
        verdict = "у допуску" if in_range else "поза допуском"
        line = [row, reference_value, actual, round(deviation * 100, 1)]
        alt_value = alternative.get(row)
        if alt_column:
            line.append(alt_value)
            if not in_range and alt_value is not None:
                alt_deviation = (float(alt_value) - float(reference_value)) / float(
                    reference_value
                )
                if abs(alt_deviation) <= TOLERANCE:
                    verdict = f"еталон відповідає колонці {alt_column}"
                    matched_any += 1
        matched_any += in_range
        line.append(verdict)
        sheet.append(line)

    sheet.append([])
    sheet.append(["У допуску ±10%", f"{matched} з {checked}"])
    if alt_column:
        sheet.append(
            ["З урахуванням другого трактування", f"{matched_any} з {checked}"]
        )
        sheet.append([])
        sheet.append(
            [
                "Пояснення: у завданні сказано «пряма відстань», тому основна "
                "колонка це велике коло за координатами Wikidata."
            ]
        )
        sheet.append(
            [
                "Рядки, де еталон розходиться, відповідають відстані дорогами "
                "(перевірено маршрутизатором OSRM) або не збігаються з жодним "
                "визначенням (Берлін-Відень 779 при прямій 524 і дорозі 680)."
            ]
        )
    for index, width in enumerate([8, 12, 16, 15, 16, 46], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    workbook.save(path)
    workbook.close()


def main() -> int:
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    missing = [name for name in FILES if not (OUTPUT / name).exists()]
    if missing:
        print(
            "Немає результатів: " + ", ".join(missing) + ". "
            "Спочатку виконайте py main.py capitals і py main.py mountains."
        )
        return 1

    for name in FILES:
        stem, suffix = name.rsplit(".", 1)
        target = SUBMISSION / f"{stem}_{SURNAME}.{suffix}"
        shutil.copyfile(OUTPUT / name, target)
        add_comparison_sheet(target, REFERENCE / name, COLUMNS[name])
        print("готово:", target.name)

    notebook = ROOT / "notebooks" / f"{SURNAME}_демонстрація.ipynb"
    print(
        "\nНа платформу здаємо: обидва файли з Data/Submission, "
        f"ноутбук {notebook.name} і архів репозиторію."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
