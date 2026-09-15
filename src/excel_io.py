"""Читання і запис книги Excel.

Запис іде через openpyxl поверх копії оригіналу. pandas тут свідомо не
використовується: `to_excel` перезаписує книгу цілком і втрачає
форматування, а вимога завдання це збереження оригінальних даних без змін.
"""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font

from .models import EnrichmentReport, RowResult

LOG_SHEET = "enrichment_log"

LOG_HEADERS = [
    "row_index",
    "target_column",
    "value",
    "unit",
    "tool",
    "source_url",
    "query",
    "confidence_level",
    "status",
    "message",
    "duration_ms",
    "llm_calls",
]


class ExcelTable:
    """Таблиця першого аркуша: заголовки плюс рядки як словники."""

    def __init__(self, path: Path | str) -> None:
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"Вхідний файл не знайдено: {self.path}")
        workbook = load_workbook(self.path)
        sheet = workbook.worksheets[0]
        self.headers: list[str] = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in sheet[1]
        ]
        self.rows: list[dict[str, Any]] = []
        for excel_row in range(2, sheet.max_row + 1):
            values = {
                header: sheet.cell(row=excel_row, column=index + 1).value
                for index, header in enumerate(self.headers)
                if header
            }
            if all(value is None for value in values.values()):
                continue  # хвостові порожні рядки книги не рахуємо
            values["__row__"] = excel_row
            self.rows.append(values)
        workbook.close()

    def sample(self, count: int = 3) -> list[dict[str, Any]]:
        """Кілька рядків для планувальника, без службового ключа."""
        return [
            {key: value for key, value in row.items() if key != "__row__"}
            for row in self.rows[:count]
        ]


def write_enriched(
    source_path: Path | str,
    output_path: Path | str,
    report: EnrichmentReport,
    target_column: str,
    extra_columns: list[str],
) -> Path:
    """Пише значення в копію оригіналу і додає аркуш з логом.

    Заповнюються тільки порожні комірки: наявні дані не чіпаються ніколи.
    Фізичний запис виконується одним потоком після пулу, бо openpyxl не
    є потокобезпечним.
    """
    source_path, output_path = Path(source_path), Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_path, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook.worksheets[0]
    headers = [
        str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]
    ]

    column_index = _ensure_column(sheet, headers, target_column)
    extra_index = {name: _ensure_column(sheet, headers, name) for name in extra_columns}

    for row_result in report.rows:
        if row_result.value is not None:
            cell = sheet.cell(row=row_result.row_index, column=column_index)
            if cell.value in (None, ""):
                cell.value = row_result.value
        for name, value in row_result.extra_values.items():
            if value is None or name not in extra_index:
                continue
            cell = sheet.cell(row=row_result.row_index, column=extra_index[name])
            if cell.value in (None, ""):
                cell.value = value

    _write_log_sheet(workbook, report.rows)
    workbook.save(output_path)
    workbook.close()
    return output_path


def _ensure_column(sheet, headers: list[str], name: str) -> int:
    """Повертає номер колонки, створюючи її, якщо такої ще немає."""
    if name in headers:
        return headers.index(name) + 1
    index = len(headers) + 1
    sheet.cell(row=1, column=index).value = name
    sheet.cell(row=1, column=index).font = Font(bold=True)
    headers.append(name)
    return index


def _write_log_sheet(workbook, rows: list[RowResult]) -> None:
    """Другий аркуш із діагностикою по кожному рядку."""
    if LOG_SHEET in workbook.sheetnames:
        del workbook[LOG_SHEET]
    sheet = workbook.create_sheet(LOG_SHEET)
    sheet.append(LOG_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(
            [
                row.row_index,
                row.target_column,
                row.value,
                row.unit,
                row.tool,
                row.source_url,
                row.query,
                row.confidence_level,
                row.status,
                row.message,
                row.duration_ms,
                row.llm_calls,
            ]
        )
    widths = [10, 16, 12, 8, 20, 46, 40, 8, 14, 60, 12, 10]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
