"""Читання і запис книги.

Головна вимога завдання: оригінальні дані лишаються без змін. Саме це
перевіряють ці тести, а не лише те, що файл узагалі записався.
"""

from __future__ import annotations

from openpyxl import load_workbook

from src.excel_io import LOG_SHEET, ExcelTable, write_enriched
from src.models import EnrichmentReport, RowResult


def test_читає_заголовки_і_рядки(mountains_file):
    table = ExcelTable(mountains_file)
    assert table.headers == ["Mountain", "Country", "height"]
    assert len(table.rows) == 3
    assert table.rows[0]["Mountain"] == "Mount Everest"
    assert table.rows[0]["__row__"] == 2


def test_приклади_рядків_без_службових_полів(mountains_file):
    sample = ExcelTable(mountains_file).sample(2)
    assert len(sample) == 2
    assert all("__row__" not in row for row in sample)


def _report(rows: list[RowResult]) -> EnrichmentReport:
    return EnrichmentReport(file_path="", rows=rows)


def test_заповнює_лише_порожні_комірки(mountains_file, tmp_path):
    """K2 у файлі вже має 8000, система не має його чіпати."""
    report = _report(
        [
            RowResult(row_index=2, target_column="height", value=8848.86, status="ok"),
            RowResult(row_index=3, target_column="height", value=9999, status="ok"),
            RowResult(row_index=4, target_column="height", value=8516, status="ok"),
        ]
    )
    output = tmp_path / "out.xlsx"
    write_enriched(mountains_file, output, report, "height", [])

    sheet = load_workbook(output).worksheets[0]
    assert sheet.cell(row=2, column=3).value == 8848.86
    assert sheet.cell(row=3, column=3).value == 8000  # не перезаписано
    assert sheet.cell(row=4, column=3).value == 8516


def test_структура_оригіналу_збережена(mountains_file, tmp_path):
    output = tmp_path / "out.xlsx"
    write_enriched(mountains_file, output, _report([]), "height", [])

    before = load_workbook(mountains_file).worksheets[0]
    after = load_workbook(output).worksheets[0]
    for row in range(1, before.max_row + 1):
        for column in range(1, before.max_column + 1):
            assert (
                before.cell(row=row, column=column).value
                == after.cell(row=row, column=column).value
            )


def test_нова_колонка_створюється_під_додаткове_трактування(mountains_file, tmp_path):
    report = _report(
        [
            RowResult(
                row_index=2,
                target_column="height",
                value=8848.86,
                status="ok",
                extra_values={"height_feet": 29031},
            )
        ]
    )
    output = tmp_path / "out.xlsx"
    write_enriched(mountains_file, output, report, "height", ["height_feet"])

    sheet = load_workbook(output).worksheets[0]
    assert sheet.cell(row=1, column=4).value == "height_feet"
    assert sheet.cell(row=2, column=4).value == 29031


def test_лог_іде_окремим_аркушем(mountains_file, tmp_path):
    report = _report(
        [
            RowResult(
                row_index=2,
                target_column="height",
                value=8848.86,
                status="ok",
                tool="wikidata_lookup",
                source_url="http://www.wikidata.org/entity/Q513",
                message="перевірка",
            )
        ]
    )
    output = tmp_path / "out.xlsx"
    write_enriched(mountains_file, output, report, "height", [])

    workbook = load_workbook(output)
    assert LOG_SHEET in workbook.sheetnames
    log = workbook[LOG_SHEET]
    assert log.cell(row=1, column=1).value == "row_index"
    assert log.cell(row=2, column=5).value == "wikidata_lookup"
    assert "Q513" in log.cell(row=2, column=6).value
