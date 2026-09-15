"""Наскрізний прогін без мережі і без моделі.

Wikidata підмінена заглушкою, тому тест перевіряє склейку всіх шарів:
план, пакетне завантаження, граф, пропуск заповнених комірок, запис
книги і підсумковий звіт.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from src import pipeline
from src.cache import Cache
from src.planner import PlanError
from src.tools import wikidata

HEIGHTS = {
    "Mount Everest": {"value": "8848.86", "item": "http://www.wikidata.org/entity/Q513"},
    "K2": {"value": "8611", "item": "http://www.wikidata.org/entity/Q43512"},
    "Lhotse": {"value": "8516", "item": "http://www.wikidata.org/entity/Q168702"},
}


@pytest.fixture
def offline(monkeypatch, tmp_path):
    """Прибирає мережу і модель: усі джерела стають локальними."""
    monkeypatch.setattr(
        wikidata,
        "lookup_property",
        lambda labels, prop, cache, type_qid=None: {
            label: HEIGHTS[label] for label in labels if label in HEIGHTS
        },
    )
    monkeypatch.setattr(pipeline, "_make_llm", lambda: None)
    monkeypatch.setattr(pipeline, "Cache", lambda: Cache(tmp_path / "cache.sqlite"))


def _plan_file(tmp_path):
    path = tmp_path / "plan.json"
    path.write_text(
        """{
  "target_column": "height",
  "value_type": "number",
  "unit": "m",
  "tool": "wikidata_lookup",
  "input_columns": {"entity": "Mountain"},
  "wikidata_property": "P2044",
  "wikidata_type": "Q8502",
  "search_query_template": "height of {Mountain}",
  "bounds": {"min_value": 100, "max_value": 9000},
  "ambiguous_interpretations": [],
  "extra_columns": {},
  "reasoning": "тест"
}""",
        encoding="utf-8",
    )
    return path


def test_наскрізний_прогін(offline, mountains_file, tmp_path):
    output = tmp_path / "result.xlsx"
    report = pipeline.process_excel(
        file_path=mountains_file,
        task_description="додай висоту гір у метрах до колонки height",
        plan_path=_plan_file(tmp_path),
        output_path=output,
        show_progress=False,
    )

    assert report.rows_total == 3
    assert report.filled == 2  # третій рядок уже заповнений у файлі
    assert report.skipped == 1
    assert report.llm_calls == 0

    sheet = load_workbook(output).worksheets[0]
    assert sheet.cell(row=2, column=3).value == 8848.86
    assert sheet.cell(row=3, column=3).value == 8000  # оригінал недоторканий
    assert sheet.cell(row=4, column=3).value == 8516


def test_повторний_прогін_нічого_не_змінює(offline, mountains_file, tmp_path):
    """Ідемпотентність: заповнені комірки не переписуються."""
    plan = _plan_file(tmp_path)
    first = tmp_path / "first.xlsx"
    pipeline.process_excel(
        file_path=mountains_file,
        task_description="висота",
        plan_path=plan,
        output_path=first,
        show_progress=False,
    )
    second = tmp_path / "second.xlsx"
    report = pipeline.process_excel(
        file_path=first,
        task_description="висота",
        plan_path=plan,
        output_path=second,
        show_progress=False,
    )
    assert report.filled == 0
    assert report.skipped == 3


def test_прапорець_overwrite_переписує(offline, mountains_file, tmp_path):
    output = tmp_path / "result.xlsx"
    report = pipeline.process_excel(
        file_path=mountains_file,
        task_description="висота",
        plan_path=_plan_file(tmp_path),
        output_path=output,
        overwrite=True,
        show_progress=False,
    )
    assert report.skipped == 0


def test_відсутній_файл_дає_зрозумілу_помилку(offline, tmp_path):
    with pytest.raises(FileNotFoundError, match="не знайдено"):
        pipeline.process_excel(
            file_path=tmp_path / "немає.xlsx",
            task_description="висота",
            plan_path=_plan_file(tmp_path),
            show_progress=False,
        )


def test_план_під_чужий_файл_відхиляється(offline, mountains_file, tmp_path):
    plan = tmp_path / "wrong.json"
    plan.write_text(
        """{
  "target_column": "distance",
  "value_type": "number",
  "tool": "wikidata_lookup",
  "input_columns": {"entity": "Capital_From"},
  "wikidata_property": "P2044",
  "search_query_template": "x {Capital_From}",
  "bounds": {},
  "ambiguous_interpretations": [],
  "extra_columns": {}
}""",
        encoding="utf-8",
    )
    with pytest.raises(PlanError, match="якої немає у файлі"):
        pipeline.process_excel(
            file_path=mountains_file,
            task_description="відстань",
            plan_path=plan,
            show_progress=False,
        )


def test_збій_пакетного_завантаження_не_валить_прогін(
    offline, monkeypatch, mountains_file, tmp_path
):
    """Пакетний запит це оптимізація, а не обов'язковий крок."""

    def explode(*args, **kwargs):
        raise RuntimeError("джерело недоступне")

    monkeypatch.setattr(wikidata, "lookup_property", explode)
    report = pipeline.process_excel(
        file_path=mountains_file,
        task_description="висота",
        plan_path=_plan_file(tmp_path),
        output_path=tmp_path / "result.xlsx",
        show_progress=False,
    )
    assert report.rows_total == 3
    assert report.filled == 0  # значень немає, але файл збережено
